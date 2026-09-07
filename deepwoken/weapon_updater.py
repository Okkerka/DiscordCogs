"""Pure mapping and conservative reconciliation for Deepwoken weapon data."""

from __future__ import annotations

from collections.abc import Callable, Collection, Mapping, Sequence
from dataclasses import dataclass
import os
from pathlib import Path
import re
import shutil
import tempfile
from typing import TypeAlias
import unicodedata

from openpyxl import Workbook

from .weapon_source import ParsedWeapon


CellValue: TypeAlias = str | int | float | None
WeaponRow: TypeAlias = dict[str, CellValue]

WORKBOOK_HEADERS = (
    "Weapon Class",
    "Weapon Type",
    "Tags",
    "Name",
    "Requirements",
    "Base Damage",
    "Scaling",
    "Armor Penetration",
    "Chip Damage",
    "Posture Damage",
    "Range",
    "Swing Speed",
    "Endlag",
    "Scaled Damage",
    "Bleed Listed",
)
CONFLICT_FIELDS = (
    "Requirements",
    "Base Damage",
    "Scaling",
    "Swing Speed",
    "Endlag",
)

_SOURCE_COLUMNS = WORKBOOK_HEADERS[3:14]
_WHITESPACE = re.compile(r"\s+")
_BLEED_LISTED = re.compile(r"\(\s*\+?(\d+(?:\.\d+)?)\s*BLD\s*\)", re.IGNORECASE)
_BLEED_MARKER = re.compile(r"\(\s*Bleed\s*\)", re.IGNORECASE)
_SOURCE_CLASS_PRIORITY = {
    "Hybrid": 0,
    "Heavy": 1,
    "Medium": 2,
    "Light": 3,
    "Exclusive": 4,
    "Offhand": 5,
    "Elemental": 6,
}
_TAG_ORDER = ("Crazy Slots", "Hybrid", "Elemental", "Offhand", "Bleed")
_NUMBER = re.compile(r"[-+]?\d+(?:\.\d+)?")
_SCALING = re.compile(r"([A-Z]{2,4})\s*:\s*(\d+(?:\.\d+)?)", re.IGNORECASE)
_REQUIREMENTS = re.compile(r"(\d+)\s*([A-Z]{2,4})", re.IGNORECASE)
_XML_ILLEGAL_CONTROL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
_NON_RANKABLE_CLASSES = frozenset({
    "special / other",
    "special/other",
    "elemental",
    "fighting style",
})
_ERROR_SAMPLE_SIZE = 3


class CandidateValidationError(ValueError):
    """Raised when a generated workbook candidate is unsafe to install."""


@dataclass(frozen=True, slots=True)
class ReconcileResult:
    rows: tuple[WeaponRow, ...]
    fetched_count: int
    source_unique_count: int
    added: tuple[str, ...]
    changed: tuple[str, ...]
    conflicts: tuple[str, ...]
    removals_retained: tuple[str, ...]


def validate_candidate(
    rows: Sequence[WeaponRow],
    *,
    active_unique_count: int,
    source_unique_count: int,
    recognized_stats: Collection[str],
    ranking_exclusions: Collection[str],
) -> None:
    """Reject incomplete or unusable rows before a candidate workbook is written.

    Validation deliberately works from copied, normalized values so a rejected
    candidate cannot alter the reconciliation result owned by its caller.
    """
    if active_unique_count < 0 or source_unique_count < 0:
        raise CandidateValidationError("Weapon counts cannot be negative.")
    if source_unique_count * 5 < active_unique_count * 4:
        raise CandidateValidationError("Weapon source is less than 80% complete.")

    validate_workbook_semantics(
        rows,
        recognized_stats=recognized_stats,
        ranking_exclusions=ranking_exclusions,
    )


def validate_workbook_semantics(
    rows: Sequence[WeaponRow],
    *,
    recognized_stats: Collection[str],
    ranking_exclusions: Collection[str],
) -> None:
    """Reject rows that are unsafe or unusable regardless of their origin."""

    if not rows:
        raise CandidateValidationError("Candidate contains no weapon rows.")

    recognized = {normalize_name(stat) for stat in recognized_stats}
    exclusions = {normalize_name(name) for name in ranking_exclusions}
    normalized_rows: list[dict[str, str]] = []
    missing_headers: list[str] = []
    unexpected_headers: list[str] = []
    has_xml_illegal_control = False
    has_formula_leading_value = False
    for row in rows:
        if not isinstance(row, Mapping):
            raise CandidateValidationError("Candidate contains an invalid row.")
        missing = [header for header in WORKBOOK_HEADERS if header not in row]
        unexpected = [str(header) for header in row if header not in WORKBOOK_HEADERS]
        if missing:
            missing_headers.extend(missing)
        if unexpected:
            unexpected_headers.extend(unexpected)
        normalized = {header: _display_text(row.get(header)) for header in WORKBOOK_HEADERS}
        has_xml_illegal_control = has_xml_illegal_control or any(
            _XML_ILLEGAL_CONTROL.search(value) for value in normalized.values()
        )
        has_formula_leading_value = has_formula_leading_value or any(
            value.startswith("=") for value in normalized.values()
        )
        normalized_rows.append(normalized)
    if missing_headers:
        raise CandidateValidationError("Candidate is missing required workbook headers.")
    if unexpected_headers:
        raise CandidateValidationError("Candidate has unexpected workbook headers.")
    if has_xml_illegal_control:
        raise CandidateValidationError("Candidate contains XML-illegal control characters.")
    if has_formula_leading_value:
        raise CandidateValidationError("Candidate contains formula-leading workbook values.")

    names: dict[str, str] = {}
    duplicate_names: list[str] = []
    blank_names = False
    for row in normalized_rows:
        display_name = _display_text(row["Name"])
        name_key = normalize_name(display_name)
        if not name_key:
            blank_names = True
            continue
        if name_key in names:
            duplicate_names.append(display_name)
        else:
            names[name_key] = display_name
    if blank_names:
        raise CandidateValidationError("Candidate contains a blank weapon name.")
    if duplicate_names:
        raise CandidateValidationError(
            "Candidate contains duplicate weapon names: " + _name_sample(duplicate_names) + "."
        )

    malformed_damage: list[str] = []
    invalid_speed: list[str] = []
    missing_scaling: list[str] = []
    unrecognized_scaling: list[str] = []
    for row in normalized_rows:
        name = _display_text(row["Name"])
        if not _is_lookup_only_offhand(row):
            if _number(row["Base Damage"]) is None:
                malformed_damage.append(name)
            speed = _number(row["Swing Speed"])
            if speed is None or speed <= 0:
                invalid_speed.append(name)
        if not _is_rankable_candidate(row, exclusions, recognized):
            continue
        scaling_stats = _scaling_stats(row["Scaling"])
        if not scaling_stats:
            missing_scaling.append(name)
        elif any(stat not in recognized for stat in scaling_stats):
            unrecognized_scaling.append(name)
    if malformed_damage:
        raise CandidateValidationError("Candidate has malformed base damage: " + _name_sample(malformed_damage) + ".")
    if invalid_speed:
        raise CandidateValidationError("Candidate has non-positive swing speed: " + _name_sample(invalid_speed) + ".")
    if missing_scaling:
        raise CandidateValidationError("Rankable weapons are missing scaling stats: " + _name_sample(missing_scaling) + ".")
    if unrecognized_scaling:
        raise CandidateValidationError("Candidate has unrecognized scaling stats: " + _name_sample(unrecognized_scaling) + ".")


def write_workbook(path: Path, rows: Sequence[WeaponRow]) -> None:
    """Write rows to a new workbook using the exact runtime worksheet schema."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "All Weapons"
    sheet.append(WORKBOOK_HEADERS)
    for row in sorted(rows, key=lambda candidate: normalize_name(candidate.get("Name"))):
        next_row = sheet.max_row + 1
        for column, header in enumerate(WORKBOOK_HEADERS, start=1):
            value = row.get(header)
            cell = sheet.cell(next_row, column, value=value)
            if isinstance(value, str):
                cell.data_type = "s"
    try:
        workbook.save(Path(path))
    finally:
        workbook.close()


def install_runtime_workbook(
    candidate_path: Path,
    active_path: Path,
    backup_path: Path,
    loader: Callable[[Path], list[WeaponRow]],
    expected_rows: Sequence[WeaponRow] | None = None,
) -> list[WeaponRow]:
    """Install a validated candidate with one fixed backup and rollback on failure."""
    candidate, active, backup, data_parent = _runtime_paths(
        candidate_path,
        active_path,
        backup_path,
    )
    backup_temp: Path | None = None
    restore_temp: Path | None = None
    candidate_replacement_attempted = False
    had_active = active.exists()
    backup_installed = False
    try:
        loaded = loader(candidate)
        if expected_rows is not None and not _workbook_rows_equivalent(loaded, expected_rows):
            raise CandidateValidationError("Reloaded candidate does not match validated weapon data.")
        if had_active:
            backup_temp = _runtime_temp_path(data_parent, "backup")
            shutil.copy2(active, backup_temp)
            os.replace(backup_temp, backup)
            backup_temp = None
            backup_installed = True

        candidate_replacement_attempted = True
        os.replace(candidate, active)
        return loaded
    except Exception as error:
        restore_error: Exception | None = None
        if had_active and backup_installed:
            try:
                restore_temp = _runtime_temp_path(data_parent, "restore")
                shutil.copy2(backup, restore_temp)
                os.replace(restore_temp, active)
                restore_temp = None
            except Exception as exc:  # pragma: no cover - filesystem failure path
                restore_error = exc
        elif (
            not had_active
            and candidate_replacement_attempted
            and not candidate.exists()
            and active.is_file()
        ):
            try:
                active.unlink(missing_ok=True)
            except OSError as exc:  # pragma: no cover - filesystem failure path
                restore_error = exc
        if restore_error is not None:
            raise RuntimeError("Unable to restore the previous runtime workbook.") from error
        raise
    finally:
        _remove_exact_path(candidate)
        _remove_exact_path(backup_temp)
        _remove_exact_path(restore_temp)


def normalize_name(value: object) -> str:
    """Return the stable comparison key for a displayed weapon name."""
    return unicodedata.normalize("NFKC", str(value) if value is not None else "").strip().casefold()


def _number(value: CellValue) -> float | None:
    match = _NUMBER.search(_display_text(value))
    return float(match.group()) if match else None


def _scaling_stats(value: CellValue) -> set[str]:
    return {normalize_name(stat) for stat, _ in _SCALING.findall(_display_text(value))}


def _is_lookup_only_offhand(row: Mapping[str, CellValue]) -> bool:
    if normalize_name(row["Weapon Class"]) not in {"special / other", "special/other"}:
        return False
    tags = {normalize_name(tag) for tag in _display_text(row["Tags"]).split(";")}
    return "offhand" in tags


def _is_rankable_candidate(
    row: Mapping[str, CellValue],
    exclusions: Collection[str],
    recognized_stats: Collection[str],
) -> bool:
    if normalize_name(row["Name"]) in exclusions:
        return False
    if normalize_name(row["Weapon Class"]) in _NON_RANKABLE_CLASSES:
        return False
    return not any(
        int(amount) > 100 and normalize_name(stat) in recognized_stats
        for amount, stat in _REQUIREMENTS.findall(_display_text(row["Requirements"]))
    )


def _name_sample(names: Sequence[str]) -> str:
    return ", ".join(sorted(set(names), key=normalize_name)[:_ERROR_SAMPLE_SIZE])


def _workbook_rows_equivalent(
    loaded: Sequence[WeaponRow],
    expected: Sequence[WeaponRow],
) -> bool:
    def signature(rows: Sequence[WeaponRow]) -> tuple[tuple[str, ...], ...]:
        return tuple(sorted(
            (
                tuple(_canonical_value(row.get(header)) for header in WORKBOOK_HEADERS)
                for row in rows
            ),
            key=lambda values: (normalize_name(values[3]), values),
        ))

    return signature(loaded) == signature(expected)


def _runtime_paths(
    candidate_path: Path,
    active_path: Path,
    backup_path: Path,
) -> tuple[Path, Path, Path, Path]:
    raw_candidate = Path(candidate_path)
    raw_active = Path(active_path)
    raw_backup = Path(backup_path)
    if any(_has_symlink_component(path) for path in (raw_candidate, raw_active, raw_backup)):
        raise CandidateValidationError("Runtime workbook paths cannot contain symlinks.")
    candidate = raw_candidate.resolve(strict=False)
    active = raw_active.resolve(strict=False)
    backup = raw_backup.resolve(strict=False)
    data_parent = active.parent
    if candidate.parent != data_parent or backup.parent != data_parent:
        raise CandidateValidationError("Runtime workbook paths must share one cog-data directory.")
    if len({candidate, active, backup}) != 3:
        raise CandidateValidationError("Runtime workbook paths must be distinct.")
    if not candidate.is_file():
        raise CandidateValidationError("Runtime workbook candidate is missing.")
    return candidate, active, backup, data_parent


def _has_symlink_component(path: Path) -> bool:
    current = path.absolute()
    while True:
        if current.is_symlink():
            return True
        parent = current.parent
        if current == parent:
            return False
        current = parent


def _runtime_temp_path(data_parent: Path, purpose: str) -> Path:
    descriptor, raw_path = tempfile.mkstemp(prefix=f".weapons-{purpose}-", suffix=".xlsx", dir=data_parent)
    os.close(descriptor)
    return Path(raw_path)


def _remove_exact_path(path: Path | None) -> None:
    if path is None:
        return
    try:
        path.unlink()
    except FileNotFoundError:
        pass


def source_to_workbook_row(row: ParsedWeapon) -> WeaponRow:
    """Map one canonical source row into the workbook's fixed fifteen columns."""
    if len(row.cells) != len(_SOURCE_COLUMNS):
        raise ValueError("Parsed weapon rows must contain exactly eleven cells.")

    source_values = dict(zip(_SOURCE_COLUMNS, row.cells, strict=True))
    name = _display_text(source_values["Name"])
    scaled_damage = _display_text(source_values["Scaled Damage"])
    bleed_match = _BLEED_LISTED.search(scaled_damage)
    bleed_listed: int | float | None = None
    if bleed_match:
        numeric = bleed_match.group(1)
        bleed_listed = float(numeric) if "." in numeric else int(numeric)

    return {
        "Weapon Class": _display_text(row.source_class),
        "Weapon Type": _display_text(row.weapon_type),
        "Tags": "Bleed" if _BLEED_MARKER.search(name) else "",
        **{column: _display_text(source_values[column]) for column in _SOURCE_COLUMNS},
        "Bleed Listed": bleed_listed,
    }


def reconcile_weapons(
    parsed: Sequence[ParsedWeapon],
    active_rows: Sequence[WeaponRow],
) -> ReconcileResult:
    """Safely reconcile parsed templates against trusted active workbook rows.

    Ambiguous source statistics never replace an active row, and ambiguous new
    weapons are omitted until their source data becomes internally consistent.
    """
    grouped_sources: dict[str, list[ParsedWeapon]] = {}
    for weapon in parsed:
        key = normalize_name(weapon.cells[0])
        if key:
            grouped_sources.setdefault(key, []).append(weapon)

    grouped_active: dict[str, list[WeaponRow]] = {}
    for row in active_rows:
        key = normalize_name(row.get("Name"))
        if key:
            grouped_active.setdefault(key, []).append(dict(row))

    active_by_name = {
        key: _select_active_row(rows)
        for key, rows in grouped_active.items()
    }
    active_conflicts = {
        key
        for key, rows in grouped_active.items()
        if _active_rows_conflict(rows)
    }

    rows: list[WeaponRow] = []
    added: list[str] = []
    changed: list[str] = []
    conflicts: list[str] = []
    seen_names: set[str] = set()

    for key in sorted(grouped_sources):
        candidates = tuple(sorted(grouped_sources[key], key=_parsed_sort_key))
        active = active_by_name.get(key)
        display_name = _display_text(candidates[0].cells[0]) or _display_text(active.get("Name") if active else "")
        seen_names.add(key)

        if key in active_conflicts:
            conflicts.append(_display_text(active.get("Name")))
            rows.append(dict(active))
            continue

        if _is_elemental_only(candidates):
            if active is None:
                conflicts.append(display_name)
            else:
                retained = _with_tag(active, "Elemental")
                rows.append(retained)
                if not _rows_match(retained, active):
                    changed.append(_display_text(active.get("Name")))
            continue

        if _has_unusable_statistics(candidates):
            conflicts.append(_display_text(active.get("Name")) if active is not None else display_name)
            if active is not None:
                rows.append(dict(active))
            continue

        candidate = _reconciled_source_row(candidates)
        display_name = _display_name(candidate, active)
        if _has_statistical_conflict(candidates):
            if active is not None:
                conflicts.append(_display_text(active.get("Name")))
                rows.append(dict(active))
            else:
                conflicts.append(display_name)
            continue

        if active is None:
            rows.append(candidate)
            added.append(display_name)
        elif _rows_match(candidate, active):
            rows.append(dict(active))
        else:
            rows.append(candidate)
            changed.append(display_name)

    removals_retained: list[str] = []
    for key, active in active_by_name.items():
        if key not in seen_names:
            rows.append(dict(active))
            removals_retained.append(_display_text(active.get("Name")))
            if key in active_conflicts:
                conflicts.append(_display_text(active.get("Name")))

    rows.sort(key=lambda row: normalize_name(row.get("Name")))
    return ReconcileResult(
        rows=tuple(rows),
        fetched_count=len(parsed),
        source_unique_count=len(grouped_sources),
        added=tuple(sorted(added, key=normalize_name)),
        changed=tuple(sorted(changed, key=normalize_name)),
        conflicts=tuple(sorted(conflicts, key=normalize_name)),
        removals_retained=tuple(sorted(removals_retained, key=normalize_name)),
    )


def _reconciled_source_row(candidates: Sequence[ParsedWeapon]) -> WeaponRow:
    mapped = [(weapon, source_to_workbook_row(weapon)) for weapon in candidates]
    crazy_slots = [(weapon, row) for weapon, row in mapped if weapon.source_class == "Crazy Slots"]
    selected_weapon, selected = (crazy_slots or mapped)[0]
    primary_weapon, primary = _primary_source(mapped)
    row = dict(selected)
    row["Weapon Class"] = _weapon_class(candidates)
    row["Weapon Type"] = _weapon_type(selected_weapon, selected, primary_weapon, primary)
    row["Tags"] = _tags(candidates, mapped, row["Weapon Class"])
    return row


def _primary_source(mapped: Sequence[tuple[ParsedWeapon, WeaponRow]]) -> tuple[ParsedWeapon, WeaponRow]:
    primary = [item for item in mapped if item[0].source_class not in {"Elemental", "Hybrid", "Offhand"}]
    if primary:
        return sorted(primary, key=_primary_sort_key)[0]
    non_elemental = [item for item in mapped if item[0].source_class != "Elemental"]
    return sorted(non_elemental or list(mapped), key=_primary_sort_key)[0]


def _weapon_class(candidates: Sequence[ParsedWeapon]) -> str:
    classes = {weapon.source_class for weapon in candidates}
    if "Offhand" in classes:
        return "Special / Other"
    if "Crazy Slots" in classes:
        return "Crazy Slots"
    if "Hybrid" in classes or len(classes & {"Light", "Medium", "Heavy"}) >= 2:
        return "Hybrid"
    primary_classes = classes - {"Elemental"}
    if not primary_classes:
        return "Unknown"
    return sorted(primary_classes, key=_source_class_sort_key)[0]


def _weapon_type(
    selected_weapon: ParsedWeapon,
    selected: WeaponRow,
    primary_weapon: ParsedWeapon,
    primary: WeaponRow,
) -> str:
    if selected_weapon.source_class == "Crazy Slots" and _has_weapon_type(selected["Weapon Type"]):
        return _display_text(selected["Weapon Type"])
    if _has_weapon_type(primary["Weapon Type"]):
        return _display_text(primary["Weapon Type"])
    return _display_text(selected["Weapon Type"])


def _tags(
    candidates: Sequence[ParsedWeapon],
    mapped: Sequence[tuple[ParsedWeapon, WeaponRow]],
    weapon_class: str,
) -> str:
    classes = {weapon.source_class for weapon in candidates}
    tags: set[str] = set()
    if weapon_class == "Crazy Slots":
        tags.add("Crazy Slots")
    if weapon_class == "Hybrid":
        tags.add("Hybrid")
    if "Elemental" in classes:
        tags.add("Elemental")
    if "Offhand" in classes:
        tags.add("Offhand")
    if any(_display_text(row["Tags"]) == "Bleed" for _, row in mapped):
        tags.add("Bleed")
    return "; ".join(tag for tag in _TAG_ORDER if tag in tags)


def _has_statistical_conflict(candidates: Sequence[ParsedWeapon]) -> bool:
    mapped = [(weapon, source_to_workbook_row(weapon)) for weapon in candidates]
    statistical = [item for item in mapped if item[0].source_class != "Offhand"] or mapped
    for field in CONFLICT_FIELDS:
        field_rows = statistical
        if field == "Requirements":
            crazy_slots = [item for item in mapped if item[0].source_class == "Crazy Slots"]
            if crazy_slots:
                field_rows = crazy_slots
        values = {_canonical_value(row[field]) for _, row in field_rows}
        if len(values) > 1:
            return True
    return False


def _has_unusable_statistics(candidates: Sequence[ParsedWeapon]) -> bool:
    for candidate in (row for row in candidates if row.source_class != "Offhand"):
        row = source_to_workbook_row(candidate)
        speed = _number(row["Swing Speed"])
        if _number(row["Base Damage"]) is None or speed is None or speed <= 0:
            return True
    return False


def _rows_match(source: WeaponRow, active: WeaponRow) -> bool:
    return all(_canonical_value(source.get(header)) == _canonical_value(active.get(header)) for header in WORKBOOK_HEADERS)


def _is_elemental_only(candidates: Sequence[ParsedWeapon]) -> bool:
    return all(weapon.source_class == "Elemental" for weapon in candidates)


def _select_active_row(rows: Sequence[WeaponRow]) -> WeaponRow:
    """Keep the most complete active row with a deterministic canonical tiebreak."""
    return dict(min(rows, key=_active_row_sort_key))


def _active_rows_conflict(rows: Sequence[WeaponRow]) -> bool:
    return len({_active_row_signature(row) for row in rows}) > 1


def _active_row_sort_key(row: WeaponRow) -> tuple[object, ...]:
    completeness = sum(bool(_canonical_value(row.get(header))) for header in WORKBOOK_HEADERS)
    return (-completeness, _active_row_signature(row))


def _active_row_signature(row: WeaponRow) -> tuple[str, ...]:
    return tuple(_canonical_value(row.get(header)) for header in WORKBOOK_HEADERS)


def _with_tag(row: WeaponRow, tag: str) -> WeaponRow:
    retained = dict(row)
    tags = {
        normalize_name(part): _display_text(part)
        for part in _display_text(row.get("Tags")).split(";")
        if _display_text(part)
    }
    tags[normalize_name(tag)] = tag
    known_tags = {normalize_name(tag) for tag in _TAG_ORDER}
    ordered = [tag for tag in _TAG_ORDER if normalize_name(tag) in tags]
    ordered.extend(tags[key] for key in sorted(tags) if key not in known_tags)
    retained["Tags"] = "; ".join(ordered)
    return retained


def _parsed_sort_key(weapon: ParsedWeapon) -> tuple[object, ...]:
    return (
        _source_class_sort_key(weapon.source_class),
        normalize_name(weapon.weapon_type),
        tuple(normalize_name(cell) for cell in weapon.cells),
        _capitalization_sort_key(weapon.cells[0]),
        tuple(_display_text(cell) for cell in weapon.cells),
        normalize_name(weapon.source_template),
    )


def _primary_sort_key(item: tuple[ParsedWeapon, WeaponRow]) -> tuple[object, ...]:
    weapon, row = item
    return (
        _source_class_sort_key(weapon.source_class),
        normalize_name(row["Weapon Type"]),
        tuple(_canonical_value(row.get(header)) for header in WORKBOOK_HEADERS[3:]),
        normalize_name(weapon.source_template),
    )


def _source_class_sort_key(source_class: object) -> tuple[int, str]:
    text = _display_text(source_class)
    return (_SOURCE_CLASS_PRIORITY.get(text, len(_SOURCE_CLASS_PRIORITY)), normalize_name(text))


def _has_weapon_type(value: CellValue) -> bool:
    return normalize_name(value) not in {"", "unknown"}


def _display_name(source: WeaponRow, active: WeaponRow | None) -> str:
    return _display_text(source.get("Name")) or _display_text(active.get("Name") if active else "")


def _display_text(value: object) -> str:
    return unicodedata.normalize("NFKC", str(value) if value is not None else "").strip()


def _canonical_value(value: object) -> str:
    return _WHITESPACE.sub(" ", _display_text(value)).strip()


def _capitalization_sort_key(value: object) -> tuple[int, str]:
    text = _display_text(value)
    return (int(bool(text) and text == text.upper()), text.casefold())
