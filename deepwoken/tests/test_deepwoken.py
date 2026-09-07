from __future__ import annotations

import asyncio
import json
import subprocess
import sys
from pathlib import Path
from tempfile import TemporaryDirectory
import textwrap
import threading
import types
import unittest
from unittest import mock
from zipfile import ZipFile

from openpyxl import Workbook
from openpyxl.utils.exceptions import IllegalCharacterError

from deepwoken.weapon_source import TemplatePayload, WeaponSourceError
from deepwoken.weapon_updater import (
    CandidateValidationError,
    WORKBOOK_HEADERS,
    install_runtime_workbook,
)


class FakeEmbed:
    def __init__(self, *, title=None, description=None, colour=None):
        self.title = title
        self.description = description
        self.colour = colour
        self.fields = []
        self.footer = None

    def add_field(self, **field):
        self.fields.append(field)

    def set_footer(self, *, text):
        self.footer = text


class FakeColour:
    @staticmethod
    def blurple():
        return 0x5865F2


class FakeSimpleMenu:
    def __init__(self, pages):
        self.pages = pages

    async def start(self, ctx):
        for page in self.pages:
            await ctx.send(embed=page)


class FakeBot:
    def __init__(self, *, owner=False):
        self.owner = owner
        self.owner_checks = 0

    async def is_owner(self, _author):
        self.owner_checks += 1
        return self.owner


class FakeContext:
    def __init__(self, *, bot=None, author=None):
        self.embeds = []
        self.messages = []
        self.bot = bot or FakeBot()
        self.author = author or object()

    async def send(self, content=None, *, embed=None):
        if embed is not None:
            self.embeds.append(embed)
        elif content is not None:
            self.messages.append(content)


class SpyLock:
    def __init__(self):
        self.entries = 0

    async def __aenter__(self):
        self.entries += 1

    async def __aexit__(self, *_args):
        return None


def passthrough_decorator(*_args, **_kwargs):
    return lambda function: function


discord = types.ModuleType("discord")
discord.Embed = FakeEmbed
discord.Colour = FakeColour

commands = types.ModuleType("redbot.core.commands")
commands.Cog = object
commands.Context = object
commands.command = passthrough_decorator
commands.is_owner = passthrough_decorator

redbot = types.ModuleType("redbot")
redbot_core = types.ModuleType("redbot.core")
redbot_core.commands = commands
redbot_data_manager = types.ModuleType("redbot.core.data_manager")
runtime_data_path = Path()


def cog_data_path(_cog):
    return runtime_data_path


redbot_data_manager.cog_data_path = cog_data_path
redbot_utils = types.ModuleType("redbot.core.utils")
redbot_menus = types.ModuleType("redbot.core.utils.menus")
redbot_menus.SimpleMenu = FakeSimpleMenu

fake_modules = {
    "discord": discord,
    "redbot": redbot,
    "redbot.core": redbot_core,
    "redbot.core.commands": commands,
    "redbot.core.data_manager": redbot_data_manager,
    "redbot.core.utils": redbot_utils,
    "redbot.core.utils.menus": redbot_menus,
}
original_modules = {name: sys.modules.get(name) for name in fake_modules}
sys.modules.update(fake_modules)

from deepwoken.deepwoken import Deepwoken, STAT_ALIASES

for name, module in original_modules.items():
    if module is None:
        sys.modules.pop(name, None)
    else:
        sys.modules[name] = module


class DeepwokenTests(unittest.IsolatedAsyncioTestCase):
    def setUp(self):
        self.runtime_directory = TemporaryDirectory()
        global runtime_data_path
        runtime_data_path = Path(self.runtime_directory.name) / "cog-data"
        self.cog = Deepwoken.__new__(Deepwoken)

    def tearDown(self):
        self.runtime_directory.cleanup()

    @staticmethod
    def weapon(name, weapon_class, requirements, scaling, base_damage):
        return {
            "Name": name,
            "Weapon Class": weapon_class,
            "Weapon Type": f"{weapon_class} Weapon",
            "Requirements": requirements,
            "Base Damage": str(base_damage),
            "Scaling": scaling,
            "Swing Speed": "1x",
            "Endlag": "-",
        }

    @staticmethod
    def write_workbook(path, rows):
        path.parent.mkdir(parents=True, exist_ok=True)
        workbook = Workbook()
        try:
            sheet = workbook.active
            sheet.title = "All Weapons"
            sheet.append(WORKBOOK_HEADERS)
            for row in rows:
                sheet.append([row.get(header) for header in WORKBOOK_HEADERS])
            workbook.save(path)
        finally:
            workbook.close()

    @classmethod
    def runtime_row(cls, name, base_damage="20"):
        return cls.weapon(name, "Heavy", "0 HVY", "HVY: 5", base_damage)

    @staticmethod
    def complete_row(name, *, base_damage="20"):
        row = {header: "" for header in WORKBOOK_HEADERS}
        row.update({
            "Weapon Class": "Heavy",
            "Weapon Type": "Greatsword",
            "Name": name,
            "Requirements": "0 HVY",
            "Base Damage": base_damage,
            "Scaling": "HVY: 5",
            "Swing Speed": "1x",
            "Endlag": "-",
        })
        return row

    @staticmethod
    def source_payload_rows(rows, *, revision=123):
        header = "!Name!!Requirements!!Base Damage!!Scaling!!Armor Penetration!!Chip Damage!!Posture Damage!!Range!!Swing Speed!!Endlag!!Scaled Damage"
        category = '| colspan="11" | [[Greatswords]]'
        rendered_rows = ["| " + " || ".join(str(value) for value in row) for row in rows]
        wikitext = "{|\n|-\n" + header + "\n|-\n" + category
        for row in rendered_rows:
            wikitext += "\n|-\n" + row
        wikitext += "\n|}"
        return TemplatePayload(
            template="AllHeavyWeapons",
            weapon_class="Heavy",
            revision_id=revision,
            revision_timestamp="2026-08-27T00:00:00Z",
            wikitext=wikitext,
        )

    @classmethod
    def source_payload(cls, *names, revision=123, base_damage="20"):
        rows = [
            (
                name, "0 HVY", base_damage, "HVY: 5", "-", "-", "4", "6", "1x", "-", "30",
            )
            for name in names
        ]
        return cls.source_payload_rows(rows, revision=revision)

    def configured_cog(self, source, *, weapons=()):
        cog = Deepwoken.__new__(Deepwoken)
        cog.bot = object()
        cog.runtime_data_path = runtime_data_path
        cog.runtime_workbook_path = runtime_data_path / "weapons.xlsx"
        cog.backup_workbook_path = runtime_data_path / "weapons.backup.xlsx"
        cog.bundled_workbook_path = Path(__file__).parents[1] / "data" / "weapons.xlsx"
        cog.workbook_path = cog.bundled_workbook_path
        cog._update_lock = asyncio.Lock()
        cog._weapon_source = source
        cog.weapons = [dict(row) for row in weapons]
        return cog

    async def test_initialization_and_lookup_use_bundled_workbook_without_creating_runtime_directory(self):
        cog = Deepwoken(object())
        ctx = FakeContext()

        await cog.dwweapon(ctx, "Unknown Weapon")

        self.assertEqual(cog.workbook_path, cog.bundled_workbook_path)
        self.assertTrue(cog.bundled_workbook_path.is_file())
        self.assertFalse(cog.runtime_data_path.exists())

    def test_initialization_prefers_valid_runtime_workbook(self):
        runtime_workbook = runtime_data_path / "weapons.xlsx"
        self.write_workbook(runtime_workbook, [self.runtime_row("Runtime Blade")])

        cog = Deepwoken(object())

        self.assertEqual(cog.workbook_path, runtime_workbook)
        self.assertEqual([row["Name"] for row in cog.weapons], ["Runtime Blade"])

    def test_semantically_malformed_runtime_workbook_uses_bundled_fallback(self):
        runtime_workbook = runtime_data_path / "weapons.xlsx"
        self.write_workbook(runtime_workbook, [self.runtime_row("Broken Blade", base_damage="unknown")])

        with self.assertLogs("deepwoken.deepwoken", level="WARNING") as logs:
            cog = Deepwoken(object())

        self.assertEqual(cog.workbook_path, cog.bundled_workbook_path)
        self.assertEqual(len(logs.output), 1)
        self.assertIn("bundled fallback", logs.output[0])

    def test_header_only_runtime_workbook_uses_bundled_fallback(self):
        runtime_workbook = runtime_data_path / "weapons.xlsx"
        self.write_workbook(runtime_workbook, [])

        with self.assertLogs("deepwoken.deepwoken", level="WARNING") as logs:
            cog = Deepwoken(object())

        self.assertEqual(cog.workbook_path, cog.bundled_workbook_path)
        self.assertEqual(len(logs.output), 1)
        self.assertIn("bundled fallback", logs.output[0])

    def test_metadata_declares_python_310_floor(self):
        info_path = Path(__file__).parents[1] / "info.json"

        metadata = json.loads(info_path.read_text("utf-8"))

        self.assertEqual(metadata.get("min_python_version"), [3, 10, 0])

    def test_malformed_runtime_workbook_warns_once_and_uses_bundled_fallback(self):
        runtime_workbook = runtime_data_path / "weapons.xlsx"
        runtime_workbook.parent.mkdir(parents=True)
        runtime_workbook.write_bytes(b"not a workbook")

        with self.assertLogs("deepwoken.deepwoken", level="WARNING") as logs:
            cog = Deepwoken(object())

        self.assertEqual(cog.workbook_path, cog.bundled_workbook_path)
        self.assertEqual(len(logs.output), 1)
        self.assertIn("bundled fallback", logs.output[0])
        self.assertLess(len(logs.output[0]), 256)

    def test_malformed_runtime_content_types_xml_uses_bundled_fallback(self):
        runtime_workbook = runtime_data_path / "weapons.xlsx"
        runtime_workbook.parent.mkdir(parents=True)
        with ZipFile(runtime_workbook, "w") as archive:
            archive.writestr("[Content_Types].xml", "<Types>")

        with self.assertLogs("deepwoken.deepwoken", level="WARNING") as logs:
            cog = Deepwoken(object())

        self.assertEqual(cog.workbook_path, cog.bundled_workbook_path)
        self.assertEqual(len(logs.output), 1)
        self.assertIn("bundled fallback", logs.output[0])
        self.assertLess(len(logs.output[0]), 256)

    def test_runtime_archive_missing_required_members_uses_bundled_fallback(self):
        runtime_workbook = runtime_data_path / "weapons.xlsx"
        runtime_workbook.parent.mkdir(parents=True)
        with ZipFile(runtime_workbook, "w") as archive:
            archive.writestr(
                "[Content_Types].xml",
                '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"></Types>',
            )

        with self.assertLogs("deepwoken.deepwoken", level="WARNING") as logs:
            cog = Deepwoken(object())

        self.assertEqual(cog.workbook_path, cog.bundled_workbook_path)
        self.assertEqual(len(logs.output), 1)
        self.assertIn("bundled fallback", logs.output[0])
        self.assertLess(len(logs.output[0]), 256)

    def test_wrong_runtime_schema_uses_bundled_fallback(self):
        runtime_workbook = runtime_data_path / "weapons.xlsx"
        runtime_workbook.parent.mkdir(parents=True)
        workbook = Workbook()
        try:
            workbook.save(runtime_workbook)
        finally:
            workbook.close()

        with self.assertLogs("deepwoken.deepwoken", level="WARNING"):
            cog = Deepwoken(object())

        self.assertEqual(cog.workbook_path, cog.bundled_workbook_path)

    def test_load_weapons_deduplicates_best_row_and_closes_workbook(self):
        path = runtime_data_path / "weapons.xlsx"
        workbook = mock.MagicMock()
        sheet = mock.MagicMock()
        workbook.__getitem__.return_value = sheet
        incomplete_row = {
            "Name": "Duplicate Blade",
            "Requirements": "",
            "Scaling": "",
            "Base Damage": "",
            "Swing Speed": "",
        }
        complete_row = {
            "Name": "Duplicate Blade",
            "Requirements": "0 HVY",
            "Scaling": "HVY: 5",
            "Base Damage": "20",
            "Swing Speed": "1x",
        }
        sheet.iter_rows.side_effect = [
            iter([tuple(mock.Mock(value=header) for header in WORKBOOK_HEADERS)]),
            iter([
                tuple(incomplete_row.get(header, "") for header in WORKBOOK_HEADERS),
                tuple(complete_row.get(header, "") for header in WORKBOOK_HEADERS),
            ]),
        ]

        with mock.patch("deepwoken.deepwoken.load_workbook", return_value=workbook):
            rows = self.cog._load_weapons(path)

        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["Name"], "Duplicate Blade")
        self.assertEqual(rows[0]["Base Damage"], "20")
        self.assertEqual(rows[0]["Scaling"], "HVY: 5")
        workbook.close.assert_called_once_with()

    async def test_dwreload_repeats_runtime_first_selection(self):
        cog = Deepwoken(object())
        runtime_workbook = runtime_data_path / "weapons.xlsx"
        self.write_workbook(runtime_workbook, [self.runtime_row("Reloaded Runtime Blade")])
        ctx = FakeContext()

        await cog.dwreload(ctx)

        self.assertEqual(cog.workbook_path, runtime_workbook)
        self.assertEqual([row["Name"] for row in cog.weapons], ["Reloaded Runtime Blade"])
        self.assertEqual(ctx.messages, ["Reloaded 1 unique weapon rows."])

    def test_concatenated_requirement_keeps_known_stat(self):
        self.assertEqual(
            self.cog._requirements("50 MEDJusticar Origin"),
            {"MED": 50},
        )

    async def test_multiple_weapon_stats_share_one_order_independent_ranking(self):
        self.cog.weapons = [
            self.weapon("Heavy Only", "Heavy", "100 HVY", "HVY: 10", 20),
            self.weapon("Medium Only", "Medium", "100 MED", "MED: 10", 18),
            self.weapon("Hybrid", "Hybrid", "100 HVY100 MED", "HVY: 5MED: 5", 10),
        ]

        heavy_first = FakeContext()
        parsed = self.cog._parse_stat_query(("heavy", "100", "medium", "100", "prof", "6"))
        self.assertIsNotNone(parsed)
        await self.cog._compare(heavy_first, *parsed)

        medium_first = FakeContext()
        parsed = self.cog._parse_stat_query(("medium", "100", "heavy", "100", "prof", "6"))
        self.assertIsNotNone(parsed)
        await self.cog._compare(medium_first, *parsed)

        self.assertEqual(heavy_first.embeds[0].title, medium_first.embeds[0].title)
        self.assertEqual(heavy_first.embeds[0].description, medium_first.embeds[0].description)
        for name in ("Heavy Only", "Medium Only", "Hybrid"):
            self.assertIn(name, heavy_first.embeds[0].description)

    async def test_ranking_paginates_every_eligible_weapon(self):
        self.cog.weapons = [
            self.weapon(f"Weapon {index:02}", "Heavy", "0 HVY", "HVY: 5", 40 - index)
            for index in range(20)
        ]
        ctx = FakeContext()

        parsed = self.cog._parse_stat_query(("heavy", "100", "6"))
        self.assertIsNotNone(parsed)
        await self.cog._compare(ctx, *parsed)

        rendered = "\n".join(embed.description for embed in ctx.embeds)
        self.assertEqual(len(ctx.embeds), 2)
        for index in range(20):
            self.assertIn(f"Weapon {index:02}", rendered)

    def test_ranking_pages_fit_discord_embed_limits_with_maximum_source_names(self):
        self.cog.weapons = [
            self.complete_row(f"{index:02}-" + "W" * 20_000)
            for index in range(16)
        ]
        stats = {stat: 100 for stat in set(STAT_ALIASES.values())}

        pages = self.cog._ranking_pages(stats, 6)

        self.assertEqual(len(pages), 2)
        self.assertTrue(all(len(page.title) <= 256 for page in pages))
        self.assertTrue(all(len(page.description) <= 4096 for page in pages))
        self.assertTrue(all(len(page.footer) <= 2048 for page in pages))
        self.assertEqual(pages[0].description.count("`DPS:`"), 15)
        self.assertEqual(pages[1].description.count("`DPS:`"), 1)

    async def test_invalid_stat_query_reports_valid_ranges(self):
        self.cog.weapons = []
        invalid_queries = (
            ("heavy", "101", "prof", "6"),
            ("heavy", "-1", "prof", "6"),
            ("heavy", "1.5", "prof", "6"),
            ("heavy", "²", "6"),
            ("heavy", "100", "prof", "⁶"),
        )

        for query in invalid_queries:
            with self.subTest(query=query):
                ctx = FakeContext()
                await self.cog.dwweapon(ctx, *query)
                self.assertEqual(len(ctx.messages), 1)
                self.assertIn("0 to 100", ctx.messages[0])
                self.assertIn("0 to 6", ctx.messages[0])

    async def test_equal_dps_weapons_use_stable_name_order(self):
        self.cog.weapons = [
            self.weapon("Zulu Blade", "Heavy", "0 HVY", "HVY: 5", 20),
            self.weapon("Alpha Blade", "Heavy", "0 HVY", "HVY: 5", 20),
        ]
        ctx = FakeContext()

        parsed = self.cog._parse_stat_query(("heavy", "100", "6"))
        self.assertIsNotNone(parsed)
        await self.cog._compare(ctx, *parsed)

        description = ctx.embeds[0].description
        self.assertLess(description.index("Alpha Blade"), description.index("Zulu Blade"))

    async def test_weapon_name_starting_with_stat_alias_remains_lookupable(self):
        self.cog.weapons = [
            self.weapon("Frost Gauntlets", "Light", "0 LHT", "LHT: 5ICE: 5", 20),
        ]
        ctx = FakeContext()

        await self.cog.dwweapon(ctx, "Frost", "Gauntlets")

        self.assertEqual(ctx.messages, [])
        self.assertEqual(ctx.embeds[0].title, "Frost Gauntlets")

    async def test_updatelist_with_extra_arguments_branches_before_lookup(self):
        source = FakeWeaponSource((self.source_payload("New Blade"),))
        cog = self.configured_cog(source)
        ctx = FakeContext(bot=FakeBot(owner=True))

        await cog.dwweapon(ctx, "UPDATELIST", "unexpected")

        self.assertEqual(ctx.messages, ["Use `[p]dwweapon updatelist` with no extra arguments."])
        self.assertEqual(source.calls, 0)
        self.assertEqual(ctx.bot.owner_checks, 0)

    async def test_non_owner_update_has_no_source_or_filesystem_side_effects(self):
        source = FakeWeaponSource((self.source_payload("New Blade"),))
        cog = self.configured_cog(source)
        lock = SpyLock()
        cog._update_lock = lock
        runtime_data_path.mkdir(parents=True)
        sentinels = {
            cog.runtime_workbook_path: b"active-sentinel",
            cog.backup_workbook_path: b"backup-sentinel",
            runtime_data_path / "candidate-sentinel.xlsx": b"candidate-sentinel",
        }
        for path, value in sentinels.items():
            path.write_bytes(value)
        ctx = FakeContext(bot=FakeBot(owner=False))

        await cog.dwweapon(ctx, "updatelist")

        self.assertEqual(len(ctx.messages), 1)
        self.assertIn("owner", ctx.messages[0].casefold())
        self.assertEqual(source.calls, 0)
        self.assertEqual(ctx.bot.owner_checks, 1)
        self.assertEqual(lock.entries, 0)
        self.assertEqual({path: path.read_bytes() for path in sentinels}, sentinels)

    async def test_owner_update_installs_before_publishing_rows_and_reports_counts(self):
        old_rows = [self.complete_row("Old Blade")]
        source = FakeWeaponSource((self.source_payload("New Blade", revision=987),))
        cog = self.configured_cog(source, weapons=old_rows)
        ctx = FakeContext(bot=FakeBot(owner=True))

        real_install = install_runtime_workbook

        def install_after_check(*args):
            self.assertEqual([row["Name"] for row in cog.weapons], ["Old Blade"])
            return real_install(*args)

        with mock.patch("deepwoken.deepwoken.install_runtime_workbook", side_effect=install_after_check):
            await cog.dwweapon(ctx, "updatelist")

        self.assertEqual(len(ctx.messages), 1)
        summary = ctx.messages[0]
        for fragment in (
            "987", "Fetched: 1", "added: 1", "changed: 0",
            "conflicts retained: 0", "removals retained: 1", "active: 2",
        ):
            self.assertIn(fragment, summary)
        self.assertEqual([row["Name"] for row in cog.weapons], ["New Blade", "Old Blade"])
        self.assertEqual(cog.workbook_path, cog.runtime_workbook_path)
        self.assertTrue(cog.runtime_workbook_path.is_file())

    async def test_concurrent_owner_updates_serialize_fetch_and_install(self):
        source = SerialWeaponSource(
            self.source_payload("First Blade", revision=1),
            self.source_payload("Second Blade", revision=2),
        )
        cog = self.configured_cog(source)
        first_ctx = FakeContext(bot=FakeBot(owner=True))
        second_ctx = FakeContext(bot=FakeBot(owner=True))

        first = asyncio.create_task(cog.dwweapon(first_ctx, "updatelist"))
        for _ in range(10):
            if source.first_entered.is_set():
                break
            await asyncio.sleep(0)
        self.assertTrue(source.first_entered.is_set())
        second = asyncio.create_task(cog.dwweapon(second_ctx, "updatelist"))
        await asyncio.sleep(0)

        self.assertEqual(source.calls, 1)
        source.release_first.set()
        await asyncio.gather(first, second)

        self.assertEqual(source.max_active, 1)
        self.assertEqual([row["Name"] for row in cog.weapons], ["First Blade", "Second Blade"])
        self.assertEqual(len(first_ctx.messages), 1)
        self.assertEqual(len(second_ctx.messages), 1)

    async def test_expected_update_failures_keep_old_memory_and_active_workbook(self):
        stages = {
            "source": ("_weapon_source", WeaponSourceError("remote detail")),
            "reconcile": ("reconcile_weapons", ValueError("untrusted name / path")),
            "validation": ("validate_candidate", CandidateValidationError("untrusted name / path")),
            "workbook": ("write_workbook", OSError("private filesystem path")),
            "illegal workbook": ("write_workbook", IllegalCharacterError("remote\x00detail")),
            "installation": ("install_runtime_workbook", RuntimeError("private filesystem path")),
        }
        for stage, (target, error) in stages.items():
            with self.subTest(stage=stage):
                old_rows = [self.complete_row("Old Blade")]
                source = FakeWeaponSource((self.source_payload("New Blade"),))
                if target == "_weapon_source":
                    source = FakeWeaponSource(error)
                cog = self.configured_cog(source, weapons=old_rows)
                self.write_workbook(cog.runtime_workbook_path, old_rows)
                before = cog.runtime_workbook_path.read_bytes()
                ctx = FakeContext(bot=FakeBot(owner=True))
                patcher = (
                    mock.patch(f"deepwoken.deepwoken.{target}", side_effect=error)
                    if target != "_weapon_source"
                    else mock.patch.object(cog, "_weapon_source", source)
                )

                with self.assertLogs("deepwoken.deepwoken", level="ERROR") as logs:
                    with patcher:
                        await cog.dwweapon(ctx, "updatelist")

                self.assertEqual(len(ctx.messages), 1)
                self.assertEqual(len(logs.output), 1)
                self.assertIn("failed", ctx.messages[0].casefold())
                self.assertNotIn(str(error), ctx.messages[0])
                self.assertEqual(cog.weapons, old_rows)
                self.assertEqual(cog.runtime_workbook_path.read_bytes(), before)
                self.assertEqual(
                    sorted(path.name for path in runtime_data_path.glob("*.xlsx")),
                    ["weapons.xlsx"],
                )
                for path in runtime_data_path.glob("*.xlsx"):
                    if path != cog.runtime_workbook_path:
                        path.unlink()

    async def test_cancellation_waits_for_workbook_worker_then_cleans_candidate(self):
        old_rows = [self.complete_row("Old Blade")]
        source = FakeWeaponSource((self.source_payload("New Blade"),))
        cog = self.configured_cog(source, weapons=old_rows)
        self.write_workbook(cog.runtime_workbook_path, old_rows)
        before = cog.runtime_workbook_path.read_bytes()
        ctx = FakeContext(bot=FakeBot(owner=True))
        worker_started = threading.Event()
        release_worker = threading.Event()

        def blocked_write(_path, _rows):
            worker_started.set()
            release_worker.wait(timeout=5)

        with mock.patch("deepwoken.deepwoken.write_workbook", side_effect=blocked_write):
            update = asyncio.create_task(cog.dwweapon(ctx, "updatelist"))
            for _ in range(100):
                if worker_started.is_set():
                    break
                await asyncio.sleep(0.01)
            self.assertTrue(worker_started.is_set())
            update.cancel()
            await asyncio.sleep(0)
            self.assertFalse(update.done())
            update.cancel()
            await asyncio.sleep(0)
            self.assertFalse(update.done())
            release_worker.set()
            with self.assertRaises(asyncio.CancelledError):
                await update

        self.assertEqual(cog.weapons, old_rows)
        self.assertEqual(ctx.messages, [])
        self.assertEqual(cog.runtime_workbook_path.read_bytes(), before)
        self.assertEqual(
            sorted(path.name for path in runtime_data_path.glob("*.xlsx")),
            ["weapons.xlsx"],
        )

    async def test_cancellation_after_successful_install_publishes_installed_rows(self):
        old_rows = [self.complete_row("Old Blade")]
        source = FakeWeaponSource((self.source_payload("New Blade"),))
        cog = self.configured_cog(source, weapons=old_rows)
        self.write_workbook(cog.runtime_workbook_path, old_rows)
        ctx = FakeContext(bot=FakeBot(owner=True))
        installer_started = threading.Event()
        release_installer = threading.Event()
        real_install = install_runtime_workbook

        def blocked_install(*args):
            installer_started.set()
            release_installer.wait(timeout=5)
            return real_install(*args)

        with mock.patch("deepwoken.deepwoken.install_runtime_workbook", side_effect=blocked_install):
            update = asyncio.create_task(cog.dwweapon(ctx, "updatelist"))
            for _ in range(100):
                if installer_started.is_set():
                    break
                await asyncio.sleep(0.01)
            self.assertTrue(installer_started.is_set())
            update.cancel()
            await asyncio.sleep(0)
            self.assertFalse(update.done())
            update.cancel()
            await asyncio.sleep(0)
            self.assertFalse(update.done())
            release_installer.set()
            with self.assertRaises(asyncio.CancelledError):
                await update

        self.assertEqual([row["Name"] for row in cog.weapons], ["New Blade", "Old Blade"])
        self.assertEqual(cog.workbook_path, cog.runtime_workbook_path)
        self.assertEqual(
            [row["Name"] for row in cog._load_weapons(cog.runtime_workbook_path)],
            ["New Blade", "Old Blade"],
        )
        self.assertEqual(ctx.messages, [])
        self.assertEqual(
            sorted(path.name for path in runtime_data_path.glob("*.xlsx")),
            ["weapons.backup.xlsx", "weapons.xlsx"],
        )

    def test_terminally_cancelled_worker_does_not_spin_during_command_cancellation(self):
        script = textwrap.dedent(
            """
            import asyncio
            from pathlib import Path
            from tempfile import TemporaryDirectory
            from unittest import mock

            from deepwoken.tests.test_deepwoken import (
                Deepwoken,
                DeepwokenTests,
                FakeBot,
                FakeContext,
                FakeWeaponSource,
            )


            async def main():
                captured_workers = []
                real_create_task = asyncio.create_task

                def capture_worker(coroutine):
                    worker = real_create_task(coroutine)
                    captured_workers.append(worker)
                    return worker

                async def blocked_to_thread(*_args):
                    await asyncio.Event().wait()

                with TemporaryDirectory() as directory:
                    runtime = Path(directory) / "cog-data"
                    cog = Deepwoken.__new__(Deepwoken)
                    cog.runtime_data_path = runtime
                    cog.runtime_workbook_path = runtime / "weapons.xlsx"
                    cog.backup_workbook_path = runtime / "weapons.backup.xlsx"
                    cog.workbook_path = Path("bundled.xlsx")
                    cog._update_lock = asyncio.Lock()
                    cog._weapon_source = FakeWeaponSource((DeepwokenTests.source_payload("New Blade"),))
                    cog.weapons = []
                    ctx = FakeContext(bot=FakeBot(owner=True))

                    with (
                        mock.patch("deepwoken.deepwoken.asyncio.to_thread", blocked_to_thread),
                        mock.patch("deepwoken.deepwoken.asyncio.create_task", capture_worker),
                    ):
                        command = real_create_task(cog.dwweapon(ctx, "updatelist"))
                        while not captured_workers:
                            await asyncio.sleep(0)
                        command.cancel()
                        await asyncio.sleep(0)
                        captured_workers[0].cancel()
                        await asyncio.sleep(0)
                        try:
                            await command
                        except asyncio.CancelledError:
                            pass
                        else:
                            raise AssertionError("command cancellation did not propagate")
                        assert command.done()
                        assert captured_workers[0].cancelled()
                        assert ctx.messages == []
                        assert list(runtime.glob("*.xlsx")) == []


            asyncio.run(main())
            print("terminated")
            """
        )

        try:
            completed = subprocess.run(
                [sys.executable, "-c", script],
                cwd=Path(__file__).parents[2],
                capture_output=True,
                text=True,
                timeout=3,
                check=False,
            )
        except subprocess.TimeoutExpired:
            self.fail("command spun after its captured worker became terminally cancelled")

        self.assertEqual(completed.returncode, 0, completed.stderr)
        self.assertEqual(completed.stdout.strip(), "terminated")

    async def test_nul_source_value_fails_safely_and_preserves_active_state(self):
        old_rows = [self.complete_row("Old Blade")]
        source = FakeWeaponSource((self.source_payload("Bad\x00Blade"),))
        cog = self.configured_cog(source, weapons=old_rows)
        self.write_workbook(cog.runtime_workbook_path, old_rows)
        before = cog.runtime_workbook_path.read_bytes()
        ctx = FakeContext(bot=FakeBot(owner=True))

        with self.assertLogs("deepwoken.deepwoken", level="ERROR") as logs:
            await cog.dwweapon(ctx, "updatelist")

        self.assertEqual(ctx.messages, ["Weapon list update failed; the previous list is still active."])
        self.assertEqual(cog.weapons, old_rows)
        self.assertEqual(cog.runtime_workbook_path.read_bytes(), before)
        self.assertEqual(
            sorted(path.name for path in runtime_data_path.glob("*.xlsx")),
            ["weapons.xlsx"],
        )
        self.assertEqual(len(logs.output), 1)
        self.assertLessEqual(len(logs.output[0]), 512)
        self.assertNotIn("\x00", logs.output[0])
        self.assertNotIn("Bad", logs.output[0])

    async def test_installed_maximum_values_keep_all_lookup_outputs_within_discord_limits(self):
        def maximum_value(prefix):
            return prefix + "X" * (20_000 - len(prefix))

        names = (
            maximum_value("Unique Alpha @everyone ** "),
            maximum_value("Shared Beta @everyone ** "),
            maximum_value("Shared Gamma @everyone ** "),
        )
        rows = [
            (
                name,
                maximum_value("0 HVY "),
                maximum_value("20 "),
                maximum_value("HVY: 5 "),
                maximum_value("- "),
                maximum_value("- "),
                maximum_value("4 "),
                maximum_value("6 "),
                maximum_value("1x "),
                maximum_value("- "),
                maximum_value("30 "),
            )
            for name in names
        ]
        source = FakeWeaponSource((self.source_payload_rows(rows),))
        cog = self.configured_cog(source)
        await cog.dwweapon(FakeContext(bot=FakeBot(owner=True)), "updatelist")

        exact = FakeContext()
        substring = FakeContext()
        multiple = FakeContext()
        await cog.dwweapon(exact, names[0])
        await cog.dwweapon(substring, "Unique", "Alpha")
        await cog.dwweapon(multiple, "Shared")

        for embed in (exact.embeds[0], substring.embeds[0]):
            total = len(embed.title or "") + len(embed.description or "") + len(embed.footer or "")
            total += sum(len(field["name"]) + len(field["value"]) for field in embed.fields)
            self.assertLessEqual(len(embed.title), 256)
            self.assertTrue(all(len(field["name"]) <= 256 for field in embed.fields))
            self.assertTrue(all(len(field["value"]) <= 1024 for field in embed.fields))
            self.assertLessEqual(total, 6000)
        self.assertEqual(len(multiple.messages), 1)
        self.assertLessEqual(len(multiple.messages[0]), 2000)
        self.assertIn("Shared Beta", multiple.messages[0])
        self.assertIn("Shared Gamma", multiple.messages[0])
        self.assertNotIn("@everyone", multiple.messages[0])
        self.assertTrue(all("@everyone" not in embed.title for embed in exact.embeds + substring.embeds))

    async def test_success_summary_bounds_conflict_names(self):
        names = tuple(f"Conflict-{index}-" + "Z" * 500 for index in range(20))
        old_rows = [self.complete_row(name, base_damage="19") for name in names]
        source = FakeWeaponSource((
            self.source_payload(*names, revision=10, base_damage="20"),
            self.source_payload(*names, revision=11, base_damage="21"),
        ))
        cog = self.configured_cog(source, weapons=old_rows)
        ctx = FakeContext(bot=FakeBot(owner=True))

        await cog.dwweapon(ctx, "updatelist")

        self.assertEqual(len(ctx.messages), 1)
        self.assertLessEqual(len(ctx.messages[0]), 2000)
        self.assertIn("conflicts retained: 20", ctx.messages[0])
        self.assertIn("+17 more", ctx.messages[0])

    async def test_successful_update_immediately_serves_lookup_and_public_ranking(self):
        source = FakeWeaponSource((self.source_payload("New Blade"),))
        cog = self.configured_cog(source)
        await cog.dwweapon(FakeContext(bot=FakeBot(owner=True)), "updatelist")

        lookup = FakeContext(bot=FakeBot(owner=False))
        ranking = FakeContext(bot=FakeBot(owner=False))
        await cog.dwweapon(lookup, "New", "Blade")
        await cog.dwweapon(ranking, "heavy", "100", "medium", "100", "prof", "6")

        self.assertEqual(lookup.embeds[0].title, "New Blade")
        self.assertIn("New Blade", ranking.embeds[0].description)

    async def test_non_scaling_sovereign_bangle_remains_lookupable_but_not_rankable(self):
        source = FakeWeaponSource((self.source_payload_rows((
            (
                "Sovereign Bangle", "30 LHT OR Oath: Blightsurger", "30", "N/A",
                "-", "5%", "8", "6", "1.04x", "-", "30",
            ),
        )),))
        cog = self.configured_cog(source)
        owner = FakeContext(bot=FakeBot(owner=True))

        await cog.dwweapon(owner, "updatelist")

        self.assertIn("Weapon list updated.", owner.messages[0])
        self.assertEqual([row["Name"] for row in cog.weapons], ["Sovereign Bangle"])
        self.assertEqual(cog._ranking_pages({"HVY": 100}, 6), [])

    async def test_non_owner_still_uses_public_lookup_and_paginated_ranking(self):
        source = FakeWeaponSource(())
        cog = self.configured_cog(source)
        cog.weapons = [
            self.weapon(f"Public Weapon {index:02}", "Heavy", "0 HVY", "HVY: 5", 40 - index)
            for index in range(20)
        ]
        lookup = FakeContext(bot=FakeBot(owner=False))
        ranking = FakeContext(bot=FakeBot(owner=False))

        await cog.dwweapon(lookup, "Public", "Weapon", "00")
        await cog.dwweapon(ranking, "heavy", "100", "medium", "100", "prof", "6")

        self.assertEqual(lookup.embeds[0].title, "Public Weapon 00")
        self.assertEqual(len(ranking.embeds), 2)
        self.assertEqual(lookup.bot.owner_checks, 0)
        self.assertEqual(ranking.bot.owner_checks, 0)


class FakeWeaponSource:
    def __init__(self, result):
        self.result = result
        self.calls = 0

    async def fetch(self):
        self.calls += 1
        if isinstance(self.result, BaseException):
            raise self.result
        return self.result


class SerialWeaponSource:
    def __init__(self, *payloads):
        self.payloads = payloads
        self.calls = 0
        self.active = 0
        self.max_active = 0
        self.first_entered = asyncio.Event()
        self.release_first = asyncio.Event()

    async def fetch(self):
        index = self.calls
        self.calls += 1
        self.active += 1
        self.max_active = max(self.max_active, self.active)
        try:
            if index == 0:
                self.first_entered.set()
                await self.release_first.wait()
            return (self.payloads[index],)
        finally:
            self.active -= 1


if __name__ == "__main__":
    unittest.main()
