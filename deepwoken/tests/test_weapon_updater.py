from __future__ import annotations

from itertools import permutations
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from unittest import mock

from openpyxl import load_workbook

from deepwoken.weapon_source import ParsedWeapon
import deepwoken.weapon_updater as weapon_updater
from deepwoken.weapon_updater import (
    CandidateValidationError,
    WORKBOOK_HEADERS,
    install_runtime_workbook,
    normalize_name,
    reconcile_weapons,
    source_to_workbook_row,
    validate_candidate,
    validate_workbook_semantics,
    write_workbook,
)


def parsed_weapon(
    name: str,
    *,
    weapon_class: str = "Light",
    weapon_type: str = "Dagger",
    requirements: str = "40 LHT",
    base_damage: str = "16",
    scaling: str = "LHT: 5",
    armor_penetration: str = "-",
    chip_damage: str = "-",
    posture_damage: str = "4",
    weapon_range: str = "6",
    swing_speed: str = "1.2x",
    endlag: str = "-",
    scaled_damage: str = "24",
) -> ParsedWeapon:
    return ParsedWeapon(
        source_template=f"All{weapon_class}Weapons",
        source_class=weapon_class,
        weapon_type=weapon_type,
        cells=(
            name,
            requirements,
            base_damage,
            scaling,
            armor_penetration,
            chip_damage,
            posture_damage,
            weapon_range,
            swing_speed,
            endlag,
            scaled_damage,
        ),
    )


def active_row(name: str, **changes: object) -> dict[str, str | int | float | None]:
    row = source_to_workbook_row(parsed_weapon(name))
    row.update(changes)
    return row


class WeaponWorkbookMappingTests(unittest.TestCase):
    def test_source_columns_map_to_the_fifteen_workbook_headers(self):
        source = parsed_weapon(
            "Flareblood Kamas (Bleed)",
            requirements="60 LHT 30 INT",
            base_damage="14",
            scaling="LHT: 8 INT: 1.5 BLD: 3",
            armor_penetration="30%",
            chip_damage="-",
            posture_damage="4",
            weapon_range="6",
            swing_speed="1.2x",
            endlag="-",
            scaled_damage="27.1 (+4.1 BLD)",
        )

        row = source_to_workbook_row(source)

        self.assertEqual(
            WORKBOOK_HEADERS,
            (
                "Weapon Class", "Weapon Type", "Tags", "Name", "Requirements",
                "Base Damage", "Scaling", "Armor Penetration", "Chip Damage",
                "Posture Damage", "Range", "Swing Speed", "Endlag",
                "Scaled Damage", "Bleed Listed",
            ),
        )
        self.assertEqual(tuple(row), WORKBOOK_HEADERS)
        self.assertEqual(tuple(row[header] for header in WORKBOOK_HEADERS[3:14]), source.cells)
        self.assertEqual(row["Weapon Class"], "Light")
        self.assertEqual(row["Weapon Type"], "Dagger")
        self.assertEqual(row["Tags"], "Bleed")
        self.assertEqual(row["Scaled Damage"], "27.1 (+4.1 BLD)")
        self.assertEqual(row["Bleed Listed"], 4.1)


class WeaponReconciliationTests(unittest.TestCase):
    def test_exact_duplicates_collapse_by_trimmed_unicode_casefolded_name(self):
        first = parsed_weapon("  Café Blade  ", weapon_class="Medium", weapon_type="Sword")
        duplicate = parsed_weapon("CAFE\u0301 BLADE", weapon_class="Medium", weapon_type="Sword")

        result = reconcile_weapons((duplicate, first), ())

        self.assertEqual(result.fetched_count, 2)
        self.assertEqual(result.source_unique_count, 1)
        self.assertEqual(result.added, ("Café Blade",))
        self.assertEqual(tuple(row["Name"] for row in result.rows), ("Café Blade",))

    def test_light_and_heavy_membership_becomes_hybrid_regardless_of_template_order(self):
        light = parsed_weapon("Wyrmtooth", weapon_class="Light", weapon_type="Sword")
        heavy = parsed_weapon("Wyrmtooth", weapon_class="Heavy", weapon_type="Sword")

        forward = reconcile_weapons((light, heavy), ())
        reverse = reconcile_weapons((heavy, light), ())

        self.assertEqual(forward.rows, reverse.rows)
        self.assertEqual(forward.rows[0]["Weapon Class"], "Hybrid")
        self.assertEqual(forward.rows[0]["Tags"], "Hybrid")

    def test_medium_and_heavy_membership_becomes_hybrid(self):
        medium = parsed_weapon("Wyrmtooth", weapon_class="Medium", weapon_type="Sword")
        heavy = parsed_weapon("Wyrmtooth", weapon_class="Heavy", weapon_type="Sword")

        result = reconcile_weapons((medium, heavy), ())

        self.assertEqual(result.rows[0]["Weapon Class"], "Hybrid")

    def test_light_medium_and_three_primary_memberships_become_hybrid_in_every_order(self):
        light = parsed_weapon("Wyrmtooth", weapon_class="Light", weapon_type="Sword")
        medium = parsed_weapon("Wyrmtooth", weapon_class="Medium", weapon_type="Sword")
        heavy = parsed_weapon("Wyrmtooth", weapon_class="Heavy", weapon_type="Sword")

        for candidates in (permutations((light, medium)), permutations((light, medium, heavy))):
            for ordering in candidates:
                with self.subTest(ordering=tuple(row.source_class for row in ordering)):
                    result = reconcile_weapons(ordering, ())

                    self.assertEqual(result.rows[0]["Weapon Class"], "Hybrid")
                    self.assertEqual(result.rows[0]["Tags"], "Hybrid")

    def test_crazy_slots_wins_class_requirements_and_uses_primary_type_when_missing(self):
        primary = parsed_weapon("Soulwrought Dagger", weapon_class="Light", weapon_type="Dagger")
        crazy_slots = parsed_weapon(
            "Soulwrought Dagger",
            weapon_class="Crazy Slots",
            weapon_type="Unknown",
            requirements="Crazy Slots",
        )

        result = reconcile_weapons((primary, crazy_slots), ())

        row = result.rows[0]
        self.assertEqual(row["Weapon Class"], "Crazy Slots")
        self.assertEqual(row["Weapon Type"], "Dagger")
        self.assertEqual(row["Requirements"], "Crazy Slots")
        self.assertEqual(row["Tags"], "Crazy Slots")

    def test_hybrid_statistics_use_an_overlapping_primary_weapon_type(self):
        hybrid = parsed_weapon("Wyrmtooth", weapon_class="Hybrid", weapon_type="Hybrid")
        medium = parsed_weapon("Wyrmtooth", weapon_class="Medium", weapon_type="Sword")

        result = reconcile_weapons((hybrid, medium), ())

        self.assertEqual(result.rows[0]["Weapon Class"], "Hybrid")
        self.assertEqual(result.rows[0]["Weapon Type"], "Sword")
        self.assertEqual(result.rows[0]["Tags"], "Hybrid")

    def test_offhand_source_becomes_lookup_only_special_other_data(self):
        offhand = ParsedWeapon(
            source_template="AllOffhandWeapons",
            source_class="Offhand",
            weapon_type="Shield",
            cells=("Targe", "10 FTD", "", "", "", "", "4", "", "", "", ""),
        )

        result = reconcile_weapons((offhand,), ())

        self.assertEqual(result.added, ("Targe",))
        self.assertEqual(result.conflicts, ())
        self.assertEqual(result.rows[0]["Weapon Class"], "Special / Other")
        self.assertEqual(result.rows[0]["Weapon Type"], "Shield")
        self.assertEqual(result.rows[0]["Tags"], "Offhand")
        validate_candidate(
            result.rows,
            active_unique_count=0,
            source_unique_count=1,
            recognized_stats={"LHT", "FTD"},
            ranking_exclusions=(),
        )

    def test_offhand_membership_keeps_primary_statistics_without_becoming_rankable(self):
        primary = parsed_weapon(
            "Silversix",
            weapon_class="Exclusive",
            weapon_type="Gun",
            base_damage="10",
        )
        offhand = ParsedWeapon(
            source_template="AllOffhandWeapons",
            source_class="Offhand",
            weapon_type="Offhand Pistol",
            cells=("Silversix", "N/A", "10 (8)", "", "-", "-", "1", "10", "", "", ""),
        )

        result = reconcile_weapons((offhand, primary), ())

        self.assertEqual(result.conflicts, ())
        self.assertEqual(result.added, ("Silversix",))
        self.assertEqual(result.rows[0]["Weapon Class"], "Special / Other")
        self.assertEqual(result.rows[0]["Weapon Type"], "Gun")
        self.assertEqual(result.rows[0]["Tags"], "Offhand")
        self.assertEqual(result.rows[0]["Scaling"], "LHT: 5")
        self.assertEqual(result.rows[0]["Swing Speed"], "1.2x")

    def test_elemental_is_a_tag_without_replacing_primary_class(self):
        primary = parsed_weapon("Hero Blade", weapon_class="Medium", weapon_type="Sword")
        elemental = parsed_weapon("Hero Blade", weapon_class="Elemental", weapon_type="Sword")

        result = reconcile_weapons((elemental, primary), ())

        self.assertEqual(result.rows[0]["Weapon Class"], "Medium")
        self.assertEqual(result.rows[0]["Tags"], "Elemental")

    def test_elemental_only_source_preserves_active_primary_class_and_adds_tag(self):
        active = active_row("Hero Blade", **{"Weapon Class": "Medium", "Tags": "Bleed"})
        elemental = parsed_weapon("Hero Blade", weapon_class="Elemental", weapon_type="Sword")

        result = reconcile_weapons((elemental,), (active,))

        self.assertEqual(result.rows[0]["Weapon Class"], "Medium")
        self.assertEqual(result.rows[0]["Tags"], "Elemental; Bleed")
        self.assertEqual(result.changed, ("Hero Blade",))

    def test_new_elemental_only_source_is_excluded_and_reported(self):
        elemental = parsed_weapon("Hero Blade", weapon_class="Elemental", weapon_type="Sword")

        result = reconcile_weapons((elemental,), ())

        self.assertEqual(result.rows, ())
        self.assertEqual(result.conflicts, ("Hero Blade",))

    def test_conflicting_scrape_statistics_retain_active_row_and_report_once(self):
        active = active_row("Whaling Knife", **{"Base Damage": "16"})
        baseline = parsed_weapon("Whaling Knife", base_damage="16")
        conflicting = parsed_weapon("Whaling Knife", base_damage="17")

        result = reconcile_weapons((baseline, conflicting), (active,))

        self.assertEqual(result.rows, (active,))
        self.assertEqual(result.conflicts, ("Whaling Knife",))
        self.assertEqual(result.changed, ())

    def test_each_conflict_field_preserves_the_active_row(self):
        fields = {
            "Requirements": ("40 LHT", "50 LHT"),
            "Base Damage": ("16", "17"),
            "Scaling": ("LHT: 5", "LHT: 6"),
            "Swing Speed": ("1.2x", "1.1x"),
            "Endlag": ("-", "0.2s"),
        }
        for field, (baseline_value, conflicting_value) in fields.items():
            with self.subTest(field=field):
                active = active_row("Whaling Knife", **{field: baseline_value})
                baseline = parsed_weapon("Whaling Knife", **{field.lower().replace(" ", "_"): baseline_value})
                conflicting = parsed_weapon("Whaling Knife", **{field.lower().replace(" ", "_"): conflicting_value})

                result = reconcile_weapons((baseline, conflicting), (active,))

                self.assertEqual(result.rows, (active,))
                self.assertEqual(result.conflicts, ("Whaling Knife",))

    def test_conflicting_new_weapon_is_excluded_and_reported_once(self):
        first = parsed_weapon("Uncertain Blade", weapon_class="Light", base_damage="16")
        second = parsed_weapon("Uncertain Blade", weapon_class="Light", base_damage="17")

        result = reconcile_weapons((second, first), ())

        self.assertEqual(result.rows, ())
        self.assertEqual(result.conflicts, ("Uncertain Blade",))
        self.assertEqual(result.added, ())

    def test_malformed_source_base_damage_preserves_the_trusted_active_row(self):
        active = active_row("Uncertain Blade")
        unusable = parsed_weapon("Uncertain Blade", base_damage="???")

        result = reconcile_weapons((unusable,), (active,))

        self.assertEqual(result.rows, (active,))
        self.assertEqual(result.conflicts, ("Uncertain Blade",))
        self.assertEqual(result.changed, ())

    def test_missing_source_speed_excludes_a_new_row(self):
        unusable = parsed_weapon("Uncertain Blade", swing_speed="")

        result = reconcile_weapons((unusable,), ())

        self.assertEqual(result.rows, ())
        self.assertEqual(result.conflicts, ("Uncertain Blade",))
        self.assertEqual(result.added, ())

    def test_nonpositive_source_speed_excludes_a_new_row(self):
        unusable = parsed_weapon("Uncertain Blade", swing_speed="0x")

        result = reconcile_weapons((unusable,), ())

        self.assertEqual(result.rows, ())
        self.assertEqual(result.conflicts, ("Uncertain Blade",))
        self.assertEqual(result.added, ())

    def test_unambiguous_new_weapon_is_added(self):
        result = reconcile_weapons((parsed_weapon("New Blade"),), ())

        self.assertEqual(result.added, ("New Blade",))
        self.assertEqual(tuple(row["Name"] for row in result.rows), ("New Blade",))

    def test_unambiguous_changed_weapon_replaces_active_row(self):
        active = active_row("Whaling Knife", **{"Armor Penetration": "20%"})
        source = parsed_weapon("Whaling Knife", armor_penetration="25%")

        result = reconcile_weapons((source,), (active,))

        self.assertEqual(result.changed, ("Whaling Knife",))
        self.assertEqual(result.rows[0]["Armor Penetration"], "25%")

    def test_active_weapon_absent_from_payload_is_retained_as_removal(self):
        active = active_row("Old Blade")

        result = reconcile_weapons((parsed_weapon("New Blade"),), (active,))

        self.assertEqual(result.removals_retained, ("Old Blade",))
        self.assertEqual(tuple(row["Name"] for row in result.rows), ("New Blade", "Old Blade"))

    def test_rows_and_summary_are_sorted_by_normalized_name(self):
        active = active_row("Zulu Blade")
        parsed = (parsed_weapon("beta Blade"), parsed_weapon("Alpha Blade"))

        result = reconcile_weapons(parsed, (active,))

        self.assertEqual(tuple(row["Name"] for row in result.rows), ("Alpha Blade", "beta Blade", "Zulu Blade"))
        self.assertEqual(result.added, ("Alpha Blade", "beta Blade"))
        self.assertEqual(result.removals_retained, ("Zulu Blade",))

    def test_input_dictionaries_are_not_mutated_and_whitespace_only_differences_do_not_change(self):
        active = active_row("Whaling Knife", **{"Requirements": " 40   LHT ", "Tags": None})
        before = dict(active)

        result = reconcile_weapons((parsed_weapon("Whaling Knife"),), (active,))

        self.assertEqual(active, before)
        self.assertEqual(result.changed, ())
        self.assertIsNot(result.rows[0], active)

    def test_duplicate_active_rows_use_stable_best_row_and_report_the_conflict(self):
        trusted = active_row("Duplicate Blade", **{"Base Damage": "16"})
        conflicting = active_row("Duplicate Blade", **{"Base Damage": "17"})

        forward = reconcile_weapons((), (trusted, conflicting))
        reverse = reconcile_weapons((), (conflicting, trusted))

        self.assertEqual(forward, reverse)
        self.assertEqual(forward.rows, (trusted,))
        self.assertEqual(forward.conflicts, ("Duplicate Blade",))
        self.assertEqual(forward.removals_retained, ("Duplicate Blade",))

    def test_numeric_zero_is_not_canonicalized_to_blank(self):
        numeric_zero = active_row("Zero Blade", **{"Base Damage": 0})
        blank = active_row("Zero Blade", **{"Base Damage": None})
        source = parsed_weapon("Zero Blade", base_damage="0")

        numeric_result = reconcile_weapons((source,), (numeric_zero,))
        blank_result = reconcile_weapons((source,), (blank,))

        self.assertEqual(normalize_name(0), "0")
        self.assertEqual(numeric_result.changed, ())
        self.assertEqual(blank_result.changed, ("Zero Blade",))


class WeaponWorkbookValidationTests(unittest.TestCase):
    recognized_stats = frozenset({"LHT", "MED", "HVY", "STR", "FTD", "AGI", "INT", "CHA", "WLL", "FIR", "ICE", "LTN", "WND", "SDW", "BLD", "MTL"})

    def validate(self, rows, *, active_unique_count=0, source_unique_count=None, ranking_exclusions=()):
        validate_candidate(
            rows,
            active_unique_count=active_unique_count,
            source_unique_count=len(rows) if source_unique_count is None else source_unique_count,
            recognized_stats=self.recognized_stats,
            ranking_exclusions=ranking_exclusions,
        )

    def test_validation_rejects_invalid_required_row_shapes_and_statistics(self):
        cases = {
            "duplicate names": (active_row("Blade"), active_row(" blade ")),
            "blank name": (active_row(""),),
            "missing header": ({key: value for key, value in active_row("Blade").items() if key != "Scaling"},),
            "malformed base damage": (active_row("Blade", **{"Base Damage": "unknown"}),),
            "non-positive swing speed": (active_row("Blade", **{"Swing Speed": "0x"}),),
            "missing scaling": (active_row("Blade", **{"Scaling": ""}),),
            "unrecognized scaling stat": (active_row("Blade", **{"Scaling": "MYST: 5"}),),
            "XML-illegal control": (active_row("Bad\x00Blade"),),
        }

        for reason, rows in cases.items():
            with self.subTest(reason=reason):
                with self.assertRaises(CandidateValidationError):
                    self.validate(rows)

    def test_validation_uses_pre_retention_source_count_for_completeness_guard(self):
        active_rows = tuple(active_row(f"Blade {index}") for index in range(10))
        result = reconcile_weapons(
            tuple(parsed_weapon(f"Blade {index}") for index in range(7)),
            active_rows,
        )

        with self.assertRaises(CandidateValidationError):
            self.validate(
                result.rows,
                active_unique_count=10,
                source_unique_count=result.source_unique_count,
            )

    def test_non_rankable_rows_do_not_require_scaling_data(self):
        rows = (
            active_row("Special", **{"Weapon Class": "Special / Other", "Scaling": ""}),
            active_row("Elemental", **{"Weapon Class": "Elemental", "Scaling": ""}),
            active_row("Fighting", **{"Weapon Class": "Fighting Style", "Scaling": ""}),
            active_row("Excluded", **{"Scaling": ""}),
            active_row("Over Cap", **{"Requirements": "101 LHT", "Scaling": ""}),
        )

        self.validate(rows, ranking_exclusions={"Excluded"})

    def test_lookup_only_offhand_rows_may_omit_numeric_weapon_stats(self):
        offhand = active_row(
            "Targe",
            **{
                "Weapon Class": "Special / Other",
                "Tags": "Offhand",
                "Base Damage": "",
                "Scaling": "",
                "Swing Speed": "",
            },
        )

        self.validate((offhand,))

    def test_excluded_row_still_requires_numeric_base_damage(self):
        excluded = active_row("Excluded", **{"Base Damage": "unknown", "Scaling": ""})

        with self.assertRaises(CandidateValidationError):
            self.validate((excluded,), ranking_exclusions={"Excluded"})

    def test_over_cap_row_still_requires_positive_swing_speed(self):
        over_cap = active_row(
            "Over Cap",
            **{"Requirements": "101 LHT", "Scaling": "", "Swing Speed": "0x"},
        )

        with self.assertRaises(CandidateValidationError):
            self.validate((over_cap,))

    def test_special_other_without_offhand_tag_still_requires_numeric_stats(self):
        special = active_row(
            "Special",
            **{"Weapon Class": "Special / Other", "Base Damage": "", "Swing Speed": ""},
        )

        with self.assertRaises(CandidateValidationError):
            self.validate((special,))

    def test_over_cap_exemption_uses_only_recognized_requirement_stats(self):
        self.validate((active_row("Over Cap", **{"Requirements": "101 LHT", "Scaling": ""}),))

        with self.assertRaises(CandidateValidationError):
            self.validate((active_row("Unknown Requirement", **{"Requirements": "101 LVL", "Scaling": ""}),))

    def test_validation_does_not_mutate_caller_rows(self):
        row = active_row("  Blade  ", **{"Scaling": " lht: 5 "})
        original = dict(row)

        self.validate((row,))

        self.assertEqual(row, original)

    def test_validation_rejects_formula_leading_candidate_values(self):
        row = active_row("Formula Blade", **{"Requirements": "=1+1"})

        with self.assertRaises(CandidateValidationError):
            self.validate((row,))

    def test_semantic_validation_rejects_zero_weapon_rows(self):
        with self.assertRaises(CandidateValidationError):
            validate_workbook_semantics(
                (),
                recognized_stats=self.recognized_stats,
                ranking_exclusions=(),
            )


class WeaponWorkbookStorageTests(unittest.TestCase):
    def test_write_workbook_uses_exact_schema_and_round_trips_values(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "candidate.xlsx"
            zulu = active_row("Zulu Blade", **{"Tags": "Bleed", "Bleed Listed": 1.5})
            alpha = active_row("Alpha Blade", **{"Tags": "Elemental", "Base Damage": 17})

            write_workbook(path, (zulu, alpha))

            workbook = load_workbook(path, data_only=True)
            try:
                self.assertEqual(workbook.sheetnames, ["All Weapons"])
                sheet = workbook["All Weapons"]
                self.assertEqual(tuple(cell.value for cell in sheet[1]), WORKBOOK_HEADERS)
                values = list(sheet.iter_rows(min_row=2, values_only=True))
            finally:
                workbook.close()

            self.assertEqual(values, [
                tuple(alpha[header] for header in WORKBOOK_HEADERS),
                tuple(zulu[header] for header in WORKBOOK_HEADERS),
            ])

    def test_write_workbook_forces_formula_leading_text_to_literal_string(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "candidate.xlsx"
            formula = "=HYPERLINK(\"https://example.invalid\",\"weapon\")"
            row = active_row("Formula Blade", **{"Requirements": formula, "Endlag": "-"})

            write_workbook(path, (row,))

            workbook = load_workbook(path, data_only=False)
            try:
                sheet = workbook["All Weapons"]
                requirements = sheet.cell(2, WORKBOOK_HEADERS.index("Requirements") + 1)
                endlag = sheet.cell(2, WORKBOOK_HEADERS.index("Endlag") + 1)
                self.assertEqual((requirements.value, requirements.data_type), (formula, "s"))
                self.assertEqual((endlag.value, endlag.data_type), ("-", "s"))
            finally:
                workbook.close()

            reopened = self._load_rows(path)
            self.assertEqual(reopened[0]["Requirements"], formula)
            self.assertEqual(reopened[0]["Endlag"], "-")

    def test_install_rejects_semantically_changed_reload_before_replacement(self):
        with TemporaryDirectory() as directory:
            data = Path(directory) / "data"
            data.mkdir()
            active = data / "weapons.xlsx"
            backup = data / "weapons.backup.xlsx"
            candidate = data / "candidate.xlsx"
            old = active_row("Old Blade")
            expected = active_row("New Blade")
            changed = active_row("New Blade", **{"Base Damage": "999"})
            write_workbook(active, (old,))
            write_workbook(candidate, (expected,))

            with self.assertRaises(CandidateValidationError):
                install_runtime_workbook(
                    candidate,
                    active,
                    backup,
                    lambda _path: [changed],
                    expected_rows=(expected,),
                )

            self.assertEqual(self._load_rows(active)[0]["Name"], "Old Blade")
            self.assertFalse(backup.exists())
            self.assertFalse(candidate.exists())

    def test_install_rejects_parent_escape(self):
        with TemporaryDirectory() as directory:
            root = Path(directory)
            data = root / "data"
            outside = root / "outside"
            data.mkdir()
            outside.mkdir()
            active = data / "weapons.xlsx"
            backup = data / "weapons.backup.xlsx"
            escaped_candidate = outside / "candidate.xlsx"
            write_workbook(escaped_candidate, (active_row("New Blade"),))

            with self.assertRaises(CandidateValidationError):
                install_runtime_workbook(escaped_candidate, active, backup, lambda path: [])

    def test_install_rejects_parent_symlink_escape(self):
        with TemporaryDirectory() as directory:
            root = Path(directory)
            data = root / "data"
            outside = root / "outside"
            data.mkdir()
            outside.mkdir()
            active = data / "weapons.xlsx"
            backup = data / "weapons.backup.xlsx"

            linked_parent = data / "linked"
            try:
                linked_parent.symlink_to(outside, target_is_directory=True)
            except OSError as exc:
                self.skipTest(f"Directory symlinks unavailable: {exc}")
            linked_candidate = linked_parent / "candidate.xlsx"
            write_workbook(linked_candidate, (active_row("New Blade"),))

            with self.assertRaises(CandidateValidationError):
                install_runtime_workbook(linked_candidate, active, backup, lambda path: [])

    def test_install_rejects_same_parent_candidate_symlink(self):
        with TemporaryDirectory() as directory:
            data = Path(directory) / "data"
            data.mkdir()
            active = data / "weapons.xlsx"
            backup = data / "weapons.backup.xlsx"
            target = data / "candidate-target.xlsx"
            candidate = data / "candidate.xlsx"
            write_workbook(target, (active_row("New Blade"),))
            try:
                candidate.symlink_to(target.name)
            except OSError as exc:
                self.skipTest(f"File symlinks unavailable: {exc}")

            with self.assertRaises(CandidateValidationError):
                install_runtime_workbook(candidate, active, backup, self._load_rows)

            self.assertTrue(candidate.is_symlink())
            self.assertTrue(target.is_file())
            self.assertFalse(active.exists())

    def test_successful_install_replaces_active_and_keeps_immediately_previous_backup(self):
        with TemporaryDirectory() as directory:
            data = Path(directory) / "data"
            data.mkdir()
            active = data / "weapons.xlsx"
            backup = data / "weapons.backup.xlsx"
            candidate = data / "candidate.xlsx"
            write_workbook(active, (active_row("Old Blade"),))
            write_workbook(candidate, (active_row("New Blade"),))

            loaded = install_runtime_workbook(candidate, active, backup, self._load_rows)

            self.assertEqual([row["Name"] for row in loaded], ["New Blade"])
            self.assertEqual(self._load_rows(active)[0]["Name"], "New Blade")
            self.assertEqual(self._load_rows(backup)[0]["Name"], "Old Blade")
            self.assertFalse(candidate.exists())
            self.assertEqual(sorted(path.name for path in data.glob("*.xlsx")), ["weapons.backup.xlsx", "weapons.xlsx"])

    def test_first_install_keeps_only_the_active_workbook(self):
        with TemporaryDirectory() as directory:
            data = Path(directory) / "data"
            data.mkdir()
            active = data / "weapons.xlsx"
            backup = data / "weapons.backup.xlsx"
            candidate = data / "candidate.xlsx"
            write_workbook(candidate, (active_row("New Blade"),))

            install_runtime_workbook(candidate, active, backup, self._load_rows)

            self.assertEqual(self._load_rows(active)[0]["Name"], "New Blade")
            self.assertFalse(backup.exists())
            self.assertEqual(sorted(path.name for path in data.glob("*.xlsx")), ["weapons.xlsx"])

    def test_second_success_overwrites_fixed_backup_without_extra_workbooks(self):
        with TemporaryDirectory() as directory:
            data = Path(directory) / "data"
            data.mkdir()
            active = data / "weapons.xlsx"
            backup = data / "weapons.backup.xlsx"
            first_candidate = data / "first-candidate.xlsx"
            second_candidate = data / "second-candidate.xlsx"
            write_workbook(active, (active_row("Old Blade"),))
            write_workbook(first_candidate, (active_row("First Blade"),))
            install_runtime_workbook(first_candidate, active, backup, self._load_rows)
            write_workbook(second_candidate, (active_row("Second Blade"),))

            install_runtime_workbook(second_candidate, active, backup, self._load_rows)

            self.assertEqual(self._load_rows(active)[0]["Name"], "Second Blade")
            self.assertEqual(self._load_rows(backup)[0]["Name"], "First Blade")
            self.assertEqual(sorted(path.name for path in data.glob("*.xlsx")), ["weapons.backup.xlsx", "weapons.xlsx"])

    def test_loader_failure_preserves_active_and_removes_candidate_and_temps(self):
        with TemporaryDirectory() as directory:
            data = Path(directory) / "data"
            data.mkdir()
            active = data / "weapons.xlsx"
            backup = data / "weapons.backup.xlsx"
            candidate = data / "candidate.xlsx"
            write_workbook(active, (active_row("Old Blade"),))
            write_workbook(candidate, (active_row("New Blade"),))

            with self.assertRaisesRegex(RuntimeError, "loader failed"):
                install_runtime_workbook(candidate, active, backup, self._failing_loader)

            self.assertEqual(self._load_rows(active)[0]["Name"], "Old Blade")
            self.assertFalse(backup.exists())
            self.assertFalse(candidate.exists())
            self.assertEqual(list(data.glob(".*.xlsx")), [])

    def test_replacement_failure_restores_active_and_removes_candidate_and_temps(self):
        with TemporaryDirectory() as directory:
            data = Path(directory) / "data"
            data.mkdir()
            active = data / "weapons.xlsx"
            backup = data / "weapons.backup.xlsx"
            candidate = data / "candidate.xlsx"
            write_workbook(active, (active_row("Old Blade"),))
            write_workbook(candidate, (active_row("New Blade"),))
            real_replace = weapon_updater.os.replace

            def fail_candidate_replace(source, target):
                if Path(source) == candidate and Path(target) == active:
                    real_replace(source, target)
                    raise OSError("replacement failed")
                return real_replace(source, target)

            with mock.patch.object(weapon_updater.os, "replace", side_effect=fail_candidate_replace):
                with self.assertRaisesRegex(OSError, "replacement failed"):
                    install_runtime_workbook(candidate, active, backup, self._load_rows)

            self.assertEqual(self._load_rows(active)[0]["Name"], "Old Blade")
            self.assertEqual(self._load_rows(backup)[0]["Name"], "Old Blade")
            self.assertFalse(candidate.exists())
            self.assertEqual(list(data.glob(".*.xlsx")), [])

    def test_first_install_post_replacement_failure_removes_failed_active_and_temps(self):
        with TemporaryDirectory() as directory:
            data = Path(directory) / "data"
            data.mkdir()
            active = data / "weapons.xlsx"
            backup = data / "weapons.backup.xlsx"
            candidate = data / "candidate.xlsx"
            write_workbook(candidate, (active_row("New Blade"),))
            real_replace = weapon_updater.os.replace

            def replace_then_fail(source, target):
                if Path(source) == candidate and Path(target) == active:
                    real_replace(source, target)
                    raise OSError("replacement failed")
                return real_replace(source, target)

            with mock.patch.object(weapon_updater.os, "replace", side_effect=replace_then_fail):
                with self.assertRaisesRegex(OSError, "replacement failed"):
                    install_runtime_workbook(candidate, active, backup, self._load_rows)

            self.assertFalse(active.exists())
            self.assertFalse(candidate.exists())
            self.assertFalse(backup.exists())
            self.assertEqual(list(data.glob(".*.xlsx")), [])

    @staticmethod
    def _load_rows(path: Path):
        workbook = load_workbook(path, data_only=True, read_only=True)
        try:
            sheet = workbook["All Weapons"]
            headers = [cell.value for cell in next(sheet.iter_rows(max_row=1))]
            return [dict(zip(headers, values, strict=True)) for values in sheet.iter_rows(min_row=2, values_only=True)]
        finally:
            workbook.close()

    @staticmethod
    def _failing_loader(path: Path):
        raise RuntimeError("loader failed")


if __name__ == "__main__":
    unittest.main()
